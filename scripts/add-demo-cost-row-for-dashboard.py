# -*- coding: utf-8 -*-
"""
Одноразово добавляет в «Фактические затраты по ОР» строку оборудования с заметной
суммой в одном месяце — чтобы на дашборде изменился график затрат по месяцам.

Запуск: py -3 scripts/add-demo-cost-row-for-dashboard.py
Повторный запуск добавит ещё одну строку (при необходимости удалите демо-строку вручную).
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
XLSX = DATA_DIR / "Фактические затраты по ОР.xlsx"
SHEET = "Фактические затраты по ОР"

MONTHS_ORDER = [
    "Январь 2025", "Февраль 2025", "Март 2025", "Апрель 2025",
    "Май 2025", "Июнь 2025", "Июль 2025", "Август 2025",
    "Сентябрь 2025", "Октябрь 2025", "Ноябрь 2025", "Декабрь 2025",
]

DEMO_LABEL = "[ДЕМО] Контроль обновления дашборда — насосная линия K7"
# Заметное значение по сравнению с типичными строками отчёта (тысячи ₽)
DEMO_AMOUNT = 2_750_000.0
TARGET_MONTH = "Июнь 2025"


def cell(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    return s if s else None


def find_sum_columns(ws):
    rows = list(ws.iter_rows(values_only=True))
    header_row_idx = None
    month_cols: dict[str, int] = {}
    for i, r in enumerate(rows):
        vals = [cell(v) for v in r]
        if vals and vals[0] and "Подразделение" in str(vals[0]):
            header_row_idx = i
            for j, v in enumerate(vals):
                if v and any(m in str(v) for m in MONTHS_ORDER):
                    month_cols[str(v)] = j
            break
    if header_row_idx is None:
        raise RuntimeError('Не найдена строка шапки с «Подразделение»')
    sub_row = rows[header_row_idx + 1]
    sum_cols: dict[str, int] = {}
    for month, start_j in month_cols.items():
        for j in range(start_j, min(start_j + 3, len(sub_row))):
            v = cell(sub_row[j])
            if v and "Сумма" in str(v):
                sum_cols[month] = j
                break
    if TARGET_MONTH not in sum_cols:
        raise RuntimeError(f"Нет колонки «Сумма» для {TARGET_MONTH}, есть: {list(sum_cols)}")
    return sum_cols


def find_itogo_row(ws) -> int:
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is not None and str(v).strip().startswith("Итого"):
            return r
    return ws.max_row + 1


def main() -> int:
    if not XLSX.exists():
        print(f"Нет файла: {XLSX}", file=sys.stderr)
        return 1

    wb = openpyxl.load_workbook(XLSX)
    if SHEET not in wb.sheetnames:
        print(f"Нет листа «{SHEET}»", file=sys.stderr)
        return 1
    ws = wb[SHEET]
    sum_cols = find_sum_columns(ws)
    col0 = sum_cols[TARGET_MONTH]  # 0-based
    col_1based = col0 + 1

    insert_at = find_itogo_row(ws)
    ws.insert_rows(insert_at, 1)
    ws.cell(insert_at, 1, DEMO_LABEL)
    ws.cell(insert_at, col_1based, DEMO_AMOUNT)

    wb.save(XLSX)
    wb.close()
    print(f"OK: строка <<{DEMO_LABEL}>>, {TARGET_MONTH}, сумма {DEMO_AMOUNT:,.0f} RUB — вставлена перед <<Итого>> (строка {insert_at}).")
    print(f"Пересоберите JSON: npm run build:dashboard:json")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
