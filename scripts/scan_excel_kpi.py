# -*- coding: utf-8 -*-
"""Scan Cursor_test2.xlsx for strategic KPI-related strings (UTF-8 report)."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "data" / "Cursor_test2.xlsx"
OUT_JSON = ROOT / "data" / "_excel_kpi_scan.json"

# Многосимвольные фразы — ищем как подстроку
PHRASES = [
    "выработк",
    "объем выпуск",
    "объём выпуск",
    "объем производства",
    "восстановительн",
    "восстановительная стоимость",
    "балансовая стоимость",
    "СННО",
    "снно",
    "средняя наработк",
    "наработка на отказ",
    "время восстановлен",
    "среднее время восстановлен",
    "коэффициент технического использования",
    "сверхурочн",
    "переработк",
    "квалификац",
    "обучен",
    "переподготовк",
    "повышен",
    "расписан",
    "в срок",
    "план-факт",
    "доля работ",
    "потерь от",
    "неготовност",
    "упущенн",
    "внепланов",
    "профилакт",
    "регламентн мероприят",
    "по состоянию",
    "MTBF",
    "MTTR",
]

# Отдельно — КТУ только как слово (не «структура»)
WORD_PATTERNS = [
    (re.compile(r"(?<![а-яёА-ЯЁ])\bКТУ\b(?![а-яёА-ЯЁ])"), "КТУ"),
    (re.compile(r"\bКтг\b", re.I), "КТГ"),
]


def norm(s: object) -> str:
    return str(s).replace("\n", " ").strip()


def scan_workbook(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict = {"file": str(path), "sheets": wb.sheetnames, "phrase_hits": [], "word_hits": []}

    for sn in wb.sheetnames:
        ws = wb[sn]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 4000:
                break
            for j, val in enumerate(row):
                if val is None:
                    continue
                s = norm(val)
                if not s:
                    continue
                low = s.lower()
                for ph in PHRASES:
                    if ph.lower() in low:
                        out["phrase_hits"].append({"sheet": sn, "row": i + 1, "col": j + 1, "match": ph, "text": s[:200]})
                        break
                for rx, label in WORD_PATTERNS:
                    if rx.search(s):
                        out["word_hits"].append({"sheet": sn, "row": i + 1, "col": j + 1, "match": label, "text": s[:200]})

    wb.close()

    # Уникализация простая по (sheet,row,col,text)
    def dedupe(items: list) -> list:
        seen = set()
        res = []
        for it in items:
            k = (it["sheet"], it["row"], it["col"], it["text"][:80])
            if k in seen:
                continue
            seen.add(k)
            res.append(it)
        return res

    out["phrase_hits"] = dedupe(out["phrase_hits"])
    out["word_hits"] = dedupe(out["word_hits"])
    return out


def extract_cost_kinds(path: Path) -> list[str]:
    """Ищем в колонке под заголовком 'Вид затрат' непустые короткие строки (типы затрат)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    sn = "Фактические затраты по ОР"
    kinds: list[str] = []
    if sn not in wb.sheetnames:
        wb.close()
        return kinds
    ws = wb[sn]
    header_row_idx = None
    col_idx = None
    rows_list: list[list] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        rows_list.append(list(row))
        if i > 900:
            break
    wb.close()

    for i, row in enumerate(rows_list):
        cells = [norm(x) if x is not None else "" for x in row]
        if "Вид затрат" in cells:
            header_row_idx = i
            col_idx = cells.index("Вид затрат")
            break
    if header_row_idx is None or col_idx is None:
        return kinds

    seen: set[str] = set()
    for row in rows_list[header_row_idx + 1 :]:
        if col_idx >= len(row):
            continue
        v = row[col_idx]
        if v is None:
            continue
        s = norm(v)
        if not s or len(s) > 60:
            continue
        if s in seen:
            continue
        # отсечь чисто числовые
        if re.fullmatch(r"[\d\s.,]+", s):
            continue
        seen.add(s)
        kinds.append(s)
    return kinds


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        sys.exit(1)

    data = scan_workbook(path)
    data["vid_zatrat_candidates"] = extract_cost_kinds(path)[:100]

    OUT_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUT_JSON}")


if __name__ == "__main__":
    main()
