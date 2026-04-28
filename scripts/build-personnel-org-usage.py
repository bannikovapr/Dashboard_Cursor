# -*- coding: utf-8 -*-
"""
Builds a personnel-by-organization dataset from:
  data/Анализ использования персонала организация.xlsx

Output:
  data/personnel_org_usage.json
"""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl


DATA_DIR = Path(__file__).resolve().parent.parent / "data"
OUT_JSON = DATA_DIR / "personnel_org_usage.json"
FIO_PATRONYMIC_RE = re.compile(r"(ович|евич|ична|овна|евна|оглы|улы|кызы)$", re.IGNORECASE)
DATE_RE = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")


def find_source() -> Path:
    explicit = DATA_DIR / "org_personnel_usage.xlsx"
    if explicit.exists():
        return explicit

    lowered = []
    for p in DATA_DIR.glob("*.xlsx"):
        lowered.append((p, p.name.lower()))

    for p, low in lowered:
        if "организац" in low and "персонал" in low:
            return p

    for p, low in lowered:
        if "организац" in low:
            return p

    raise FileNotFoundError("Cannot locate source XLSX for personnel usage by organization.")


def clean_text(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def parse_hours(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    s = str(value).replace("\xa0", " ").strip()
    if not s:
        return 0.0

    # 1 640 -> 1640
    s_compact = re.sub(r"(?<=\d)\s+(?=\d)", "", s)
    nums = [int(x) for x in re.findall(r"\d+", s_compact)]
    if nums and re.search(r"[A-Za-zА-Яа-яЁё]", s_compact):
        h = nums[0]
        m = nums[1] if len(nums) > 1 else 0
        sec = nums[2] if len(nums) > 2 else 0
        return h + m / 60.0 + sec / 3600.0

    try:
        return float(s_compact.replace(",", "."))
    except ValueError:
        return 0.0


def is_fio(value: str | None) -> bool:
    if not value:
        return False
    parts = [x for x in str(value).split() if x]
    if len(parts) != 3:
        return False
    if any(any(ch.isdigit() for ch in p) for p in parts):
        return False
    if not FIO_PATRONYMIC_RE.search(parts[2]):
        return False
    return all(p[0].isalpha() and p[0].upper() == p[0] for p in parts)


def is_org_row(value: str, current_org: str | None) -> bool:
    if current_org is None:
        return True
    low = value.lower()
    if any(x in value for x in ('"', "«", "»")):
        return True
    if low.startswith(("рк ", "ооо", "ао", "пао", "зао", "оао", "ип", "llc", "inc", "ltd", "gmbh")):
        return True
    if any(token in low for token in ("организац", "компан", "предприят")):
        return True
    return False


def find_header_start(ws) -> int:
    for r in range(1, min(ws.max_row, 60) + 1):
        v = clean_text(ws.cell(r, 1).value)
        if not v:
            continue
        low = v.lower()
        if low == "организация" or low.startswith("организация "):
            return r
    raise RuntimeError("Cannot find header row with first column 'Организация'.")


def build_month_specs(ws, header_row: int):
    month_cols = []
    total_col = None

    for c in range(2, ws.max_column + 1):
        v = clean_text(ws.cell(header_row, c).value)
        if not v:
            continue
        if DATE_RE.match(v):
            month_cols.append((v, c))
        elif "итого" in v.lower():
            total_col = c

    specs = []
    for idx, (month_label, start_col) in enumerate(month_cols):
        end_col = month_cols[idx + 1][1] if idx + 1 < len(month_cols) else (total_col or ws.max_column + 1)
        labeled_cols = []
        for c in range(start_col, end_col):
            marker = clean_text(ws.cell(header_row + 2, c).value)
            if marker:
                labeled_cols.append(c)
        if not labeled_cols:
            continue
        fact_col = labeled_cols[0]
        plan_col = labeled_cols[1] if len(labeled_cols) > 1 else None
        specs.append({"month": month_label, "fact_col": fact_col, "plan_col": plan_col})

    total_spec = None
    if total_col:
        labeled_cols = []
        for c in range(total_col, ws.max_column + 1):
            marker = clean_text(ws.cell(header_row + 2, c).value)
            if marker:
                labeled_cols.append(c)
        if labeled_cols:
            total_spec = {
                "fact_col": labeled_cols[0],
                "plan_col": labeled_cols[1] if len(labeled_cols) > 1 else None,
            }

    return specs, total_spec


def make_month_maps(monthly_values, months):
    fact = {}
    plan = {}
    for month in months:
        x = monthly_values.get(month, {})
        fact[month] = round(float(x.get("fact_h") or 0.0), 2)
        plan[month] = round(float(x.get("plan_h") or 0.0), 2)
    return fact, plan


def build():
    src_xlsx = find_source()
    wb = openpyxl.load_workbook(src_xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row = find_header_start(ws)
    data_start_row = header_row + 3
    month_specs, total_spec = build_month_specs(ws, header_row)
    months = [x["month"] for x in month_specs]

    rows = []
    current_org = None
    current_department = None

    for r in range(data_start_row, ws.max_row + 1):
        label = clean_text(ws.cell(r, 1).value)
        if not label:
            continue
        if label.lower() == "итого":
            continue

        monthly = {}
        sum_fact = 0.0
        sum_plan = 0.0
        for spec in month_specs:
            month = spec["month"]
            fact_h = parse_hours(ws.cell(r, spec["fact_col"]).value) if spec.get("fact_col") else 0.0
            plan_h = parse_hours(ws.cell(r, spec["plan_col"]).value) if spec.get("plan_col") else 0.0
            monthly[month] = {"fact_h": round(fact_h, 2), "plan_h": round(plan_h, 2)}
            sum_fact += fact_h
            sum_plan += plan_h

        total_fact = parse_hours(ws.cell(r, total_spec["fact_col"]).value) if total_spec and total_spec.get("fact_col") else 0.0
        total_plan = parse_hours(ws.cell(r, total_spec["plan_col"]).value) if total_spec and total_spec.get("plan_col") else 0.0
        fact_h = total_fact if total_fact > 0 else sum_fact
        plan_h = total_plan if total_plan > 0 else sum_plan

        if is_fio(label):
            row_type = "employee"
        elif is_org_row(label, current_org):
            row_type = "organization"
        else:
            row_type = "department"

        if row_type == "organization":
            current_org = label
            current_department = None
        elif row_type == "department":
            current_department = label

        utilization = (fact_h / plan_h * 100.0) if plan_h > 0 else None
        rows.append(
            {
                "type": row_type,
                "label": label,
                "organization": current_org,
                "department": current_department if row_type == "employee" else (label if row_type == "department" else None),
                "employee": label if row_type == "employee" else None,
                "fact_h": round(fact_h, 2),
                "plan_h": round(plan_h, 2),
                "utilization_pct": round(utilization, 2) if utilization is not None else None,
                "monthly": monthly,
            }
        )

    wb.close()

    employee_rows = [x for x in rows if x["type"] == "employee" and x["employee"]]
    for x in employee_rows:
        monthly_fact, monthly_plan = make_month_maps(x.get("monthly") or {}, months)
        x["monthly_fact_h"] = monthly_fact
        x["monthly_plan_h"] = monthly_plan
        x.pop("monthly", None)
        x.pop("type", None)
        x.pop("label", None)

    employee_rows.sort(key=lambda x: x["fact_h"], reverse=True)

    dep_aggr = defaultdict(lambda: {"fact_h": 0.0, "plan_h": 0.0, "employees": 0, "monthly_fact_h": defaultdict(float), "monthly_plan_h": defaultdict(float)})
    org_aggr = defaultdict(lambda: {"fact_h": 0.0, "plan_h": 0.0, "employees": 0, "monthly_fact_h": defaultdict(float), "monthly_plan_h": defaultdict(float)})

    for e in employee_rows:
        dep_key = (e["organization"], e["department"])
        dep_aggr[dep_key]["fact_h"] += float(e["fact_h"] or 0.0)
        dep_aggr[dep_key]["plan_h"] += float(e["plan_h"] or 0.0)
        dep_aggr[dep_key]["employees"] += 1
        for month in months:
            dep_aggr[dep_key]["monthly_fact_h"][month] += float((e.get("monthly_fact_h") or {}).get(month) or 0.0)
            dep_aggr[dep_key]["monthly_plan_h"][month] += float((e.get("monthly_plan_h") or {}).get(month) or 0.0)

        org_key = e["organization"]
        org_aggr[org_key]["fact_h"] += float(e["fact_h"] or 0.0)
        org_aggr[org_key]["plan_h"] += float(e["plan_h"] or 0.0)
        org_aggr[org_key]["employees"] += 1
        for month in months:
            org_aggr[org_key]["monthly_fact_h"][month] += float((e.get("monthly_fact_h") or {}).get(month) or 0.0)
            org_aggr[org_key]["monthly_plan_h"][month] += float((e.get("monthly_plan_h") or {}).get(month) or 0.0)

    departments = []
    for (org, dep), v in dep_aggr.items():
        util = (v["fact_h"] / v["plan_h"] * 100.0) if v["plan_h"] > 0 else None
        departments.append(
            {
                "organization": org,
                "department": dep,
                "employees": int(v["employees"]),
                "fact_h": round(v["fact_h"], 2),
                "plan_h": round(v["plan_h"], 2),
                "utilization_pct": round(util, 2) if util is not None else None,
                "monthly_fact_h": {m: round(v["monthly_fact_h"][m], 2) for m in months},
                "monthly_plan_h": {m: round(v["monthly_plan_h"][m], 2) for m in months},
            }
        )
    departments.sort(key=lambda x: x["fact_h"], reverse=True)

    organizations = []
    for org, v in org_aggr.items():
        util = (v["fact_h"] / v["plan_h"] * 100.0) if v["plan_h"] > 0 else None
        organizations.append(
            {
                "organization": org,
                "employees": int(v["employees"]),
                "fact_h": round(v["fact_h"], 2),
                "plan_h": round(v["plan_h"], 2),
                "utilization_pct": round(util, 2) if util is not None else None,
                "monthly_fact_h": {m: round(v["monthly_fact_h"][m], 2) for m in months},
                "monthly_plan_h": {m: round(v["monthly_plan_h"][m], 2) for m in months},
            }
        )
    organizations.sort(key=lambda x: x["fact_h"], reverse=True)

    top_departments = departments[:12]
    payload = {
        "meta": {
            "source": src_xlsx.name,
            "sheet": ws.title,
            "purpose": "Personnel usage by organizations/departments/employees for dashboard and AI agent",
            "months": months,
            "organizations_count": len(organizations),
            "departments_count": len(departments),
            "employees_count": len(employee_rows),
        },
        "table": {
            "title": "Использование персонала по организациям и подразделениям",
            "columns": [
                {"key": "organization", "label": "Организация"},
                {"key": "department", "label": "Подразделение"},
                {"key": "employee", "label": "Сотрудник"},
                {"key": "fact_h", "label": "Факт, ч"},
                {"key": "plan_h", "label": "План, ч"},
                {"key": "utilization_pct", "label": "Выполнение, %"},
            ],
            "rows": employee_rows,
        },
        "departments": departments,
        "organizations": organizations,
        "chart": {
            "type": "chart",
            "title": "Топ подразделений по фактическим трудозатратам (ч)",
            "chartType": "bar",
            "categories": [f"{x['department']} · {x['organization']}" for x in top_departments],
            "series": [
                {"name": "Факт, ч", "data": [x["fact_h"] for x in top_departments]},
                {"name": "План, ч", "data": [x["plan_h"] for x in top_departments]},
            ],
        },
    }

    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Saved: {OUT_JSON}")
    print(
        json.dumps(
            {
                "source": src_xlsx.name,
                "months": len(months),
                "organizations": len(organizations),
                "departments": len(departments),
                "employees": len(employee_rows),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    build()
