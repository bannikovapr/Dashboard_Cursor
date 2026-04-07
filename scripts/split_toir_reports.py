# -*- coding: utf-8 -*-
"""
Разделяет единый Excel-файл ТОиР на 7 отдельных файлов (по одному листу).
Имена файлов совпадают с именами листов.
"""
from __future__ import annotations

from pathlib import Path
import openpyxl


SRC = Path(__file__).resolve().parent.parent / "data" / "Cursor_test2.xlsx"
OUT_DIR = Path(__file__).resolve().parent.parent / "data"

REPORT_SHEETS = [
    "Анализ отказов",
    "КТГ",
    "Наработка на отказ",
    "Простой",
    "Процент износа",
    "Список оборудования",
    "Фактические затраты по ОР",
]


def copy_sheet_values(src_ws, dst_ws):
    for row in src_ws.iter_rows():
        for cell in row:
            dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)


def main() -> int:
    if not SRC.exists():
        raise FileNotFoundError(f"Не найден исходный файл: {SRC}")

    wb = openpyxl.load_workbook(str(SRC), read_only=False, data_only=False)
    try:
        missing = [s for s in REPORT_SHEETS if s not in wb.sheetnames]
        if missing:
            raise ValueError(f"В исходном файле нет листов: {missing}")

        OUT_DIR.mkdir(parents=True, exist_ok=True)
        created = []
        for sheet_name in REPORT_SHEETS:
            src_ws = wb[sheet_name]
            out_wb = openpyxl.Workbook()
            dst_ws = out_wb.active
            dst_ws.title = sheet_name
            copy_sheet_values(src_ws, dst_ws)
            out_path = OUT_DIR / f"{sheet_name}.xlsx"
            out_wb.save(out_path)
            out_wb.close()
            created.append(out_path.name)

        print("Созданы файлы отчетов:")
        for name in created:
            print(f" - {name}")
    finally:
        wb.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
