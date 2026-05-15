# -*- coding: utf-8 -*-
"""
Разбор отчётов по персоналу для встраивания в data/toir.json (обязательны при сборке дашборда).

Источники в data/: «Использование персонала*.xlsx», «Анализ использования персонала*.xlsx»
(точные шаблоны имён — см. find_personnel_usage_xlsx / find_org_personnel_analysis_xlsx).
"""
from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

import openpyxl

FIO_PATRONYMIC_ORG_RE = re.compile(r"(ович|евич|ична|овна|евна|оглы|улы|кызы)$", re.IGNORECASE)
DATE_RE = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")

_CYR_U = r"\u0410-\u042f\u0401"
_CYR_L = r"\u0430-\u044f\u0451"
_WORD = rf"[{_CYR_U}][{_CYR_L}]+(?:-[{_CYR_U}][{_CYR_L}]+)?"
FIO_RE_USAGE = re.compile(rf"^{_WORD}\s+{_WORD}\s+{_WORD}$")

HOURS_RE_USAGE = re.compile(
    r"^\s*(?:(?P<h>\d+)\s*\u0447\.\s*)?(?:(?P<m>\d+)\s*\u043c\u0438\u043d\.\s*)?(?:(?P<s>\d+)\s*\u0441\.)?\s*$",
    re.IGNORECASE,
)


def find_personnel_usage_xlsx(data_dir: Path) -> Path | None:
    preferred = data_dir / "Использование персонала.xlsx"
    if preferred.exists():
        return preferred
    for p in data_dir.glob("*.xlsx"):
        low = p.name.lower()
        if "использован" in low and "персонал" in low:
            return p
    return None


def find_org_personnel_analysis_xlsx(data_dir: Path) -> Path | None:
    for name in (
        "Анализ использования персонала организация.xlsx",
        "Анализ использования персонала.xlsx",
    ):
        p = data_dir / name
        if p.exists():
            return p
    legacy = data_dir / "org_personnel_usage.xlsx"
    if legacy.exists():
        return legacy
    for p in data_dir.glob("*.xlsx"):
        low = p.name.lower()
        if "анализ" in low and "персонал" in low:
            return p
    for p in data_dir.glob("*.xlsx"):
        low = p.name.lower()
        if "организац" in low and "персонал" in low:
            return p
    return None


def _clean_usage(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return v


def _parse_hours_usage(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", " ")
    m = HOURS_RE_USAGE.match(s)
    if m:
        h = int(m.group("h") or 0)
        mm = int(m.group("m") or 0)
        ss = int(m.group("s") or 0)
        return h + (mm / 60.0) + (ss / 3600.0)
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return 0.0


def _is_fio_usage(text: str) -> bool:
    return bool(FIO_RE_USAGE.match(text or ""))


def _find_usage_columns(rows):
    header_main_idx = None
    for i, row in enumerate(rows):
        if _clean_usage(row[0]) == "Сотрудник":
            header_main_idx = i
            break
    if header_main_idx is None:
        raise RuntimeError("Не найдена строка заголовка с первым столбцом «Сотрудник»")

    header_main = rows[header_main_idx]
    header_kinds = rows[header_main_idx + 2] if header_main_idx + 2 < len(rows) else []

    total_start = None
    for j, v in enumerate(header_main):
        cv = _clean_usage(v)
        if isinstance(cv, str) and "итого" in cv.lower():
            total_start = j
            break
    if total_start is None:
        total_start = max(0, len(header_main) - 2)

    fact_col = None
    plan_col = None
    for j in range(total_start, len(header_kinds)):
        cv = _clean_usage(header_kinds[j])
        if cv == "Факт" and fact_col is None:
            fact_col = j
        elif cv == "План" and plan_col is None:
            plan_col = j

    if fact_col is None:
        fact_col = total_start
    if plan_col is None:
        plan_col = min(total_start + 1, max(0, len(header_main) - 1))

    return header_main_idx + 3, fact_col, plan_col


def build_personnel_usage_payload(data_dir: Path) -> dict | None:
    """Пейлоад как раньше в personnel_dlp_test.json (meta, table, chart)."""
    src = find_personnel_usage_xlsx(data_dir)
    if src is None or not src.exists():
        return None

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    sheet_title = ws.title
    rows = list(ws.iter_rows(values_only=True))

    data_start, fact_col, plan_col = _find_usage_columns(rows)

    agg = {}
    for row in rows[data_start:]:
        if not row:
            continue
        fio = _clean_usage(row[0]) if len(row) > 0 else None
        if not isinstance(fio, str) or not _is_fio_usage(fio):
            continue

        fact_h = _parse_hours_usage(row[fact_col] if fact_col < len(row) else None)
        plan_h = _parse_hours_usage(row[plan_col] if plan_col < len(row) else None)

        cur = agg.get(fio)
        if cur is None:
            agg[fio] = {"employee": fio, "fact_h": fact_h, "plan_h": plan_h}
        else:
            cur["fact_h"] += fact_h
            cur["plan_h"] += plan_h

    wb.close()

    rows_out = []
    for v in agg.values():
        plan_h = v["plan_h"]
        util = (v["fact_h"] / plan_h * 100.0) if plan_h > 0 else None
        rows_out.append(
            {
                "employee": v["employee"],
                "fact_h": round(v["fact_h"], 2),
                "plan_h": round(v["plan_h"], 2),
                "utilization_pct": round(util, 2) if util is not None else None,
            }
        )

    rows_out.sort(key=lambda x: x["fact_h"], reverse=True)
    top = rows_out[:10]

    return {
        "meta": {
            "source": src.name,
            "sheet": sheet_title,
            "purpose": "Personnel utilization dataset for dashboard assistant",
            "employees_count": len(rows_out),
        },
        "table": {
            "title": "Ремонтные работы сотрудников за год (итого)",
            "columns": [
                {"key": "employee", "label": "Сотрудник"},
                {"key": "fact_h", "label": "Факт, ч"},
                {"key": "plan_h", "label": "План, ч"},
                {"key": "utilization_pct", "label": "Выполнение, %"},
            ],
            "rows": rows_out,
        },
        "chart": {
            "type": "chart",
            "title": "Топ-10 сотрудников по выполненным ремонтным работам (факт, ч)",
            "chartType": "bar",
            "categories": [x["employee"] for x in top],
            "series": [
                {"name": "Факт, ч", "data": [x["fact_h"] for x in top]},
                {"name": "План, ч", "data": [x["plan_h"] for x in top]},
            ],
        },
    }


def _clean_org_text(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def _parse_hours_org(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    s = str(value).replace("\xa0", " ").strip()
    if not s:
        return 0.0

    s_compact = re.sub(r"(?<=\d)\s+(?=\d)", "", s)
    nums = [int(x) for x in re.findall(r"\d+", s_compact)]
    if nums and re.search(r"[A-Za-zА-Яа-яЁё]", s_compact):
        h = nums[0]
        m = nums[1] if len(nums) > 1 else 0
        sec = nums[2] if len(nums) > 2 else 0
        return h + m / 60.0 + sec / 3600.0

    try:
        return float(s_compact.replace(",", "."))
    except ValueError:
        return 0.0


def _is_fio_org(value: str | None) -> bool:
    if not value:
        return False
    parts = [x for x in str(value).split() if x]
    if len(parts) != 3:
        return False
    if any(any(ch.isdigit() for ch in p) for p in parts):
        return False
    if not FIO_PATRONYMIC_ORG_RE.search(parts[2]):
        return False
    return all(p[0].isalpha() and p[0].upper() == p[0] for p in parts)


def _is_org_row_org(value: str, current_org: str | None) -> bool:
    if current_org is None:
        return True
    low = value.lower()
    if any(x in value for x in ('"', "«", "»")):
        return True
    if low.startswith(("рк ", "ооо", "ао", "пао", "зао", "оао", "ип", "llc", "inc", "ltd", "gmbh")):
        return True
    if any(token in low for token in ("организац", "компан", "предприят")):
        return True
    return False


def _find_org_header_start(ws) -> int:
    for r in range(1, min(ws.max_row, 60) + 1):
        v = _clean_org_text(ws.cell(r, 1).value)
        if not v:
            continue
        low = v.lower()
        if low == "организация" or low.startswith("организация "):
            return r
    raise RuntimeError("Не найдена строка заголовка с первым столбцом «Организация»")


def _build_org_month_specs(ws, header_row: int):
    month_cols = []
    total_col = None

    for c in range(2, ws.max_column + 1):
        v = _clean_org_text(ws.cell(header_row, c).value)
        if not v:
            continue
        if DATE_RE.match(v):
            month_cols.append((v, c))
        elif "итого" in v.lower():
            total_col = c

    specs = []
    for idx, (month_label, start_col) in enumerate(month_cols):
        end_col = month_cols[idx + 1][1] if idx + 1 < len(month_cols) else (total_col or ws.max_column + 1)
        labeled_cols = []
        for c in range(start_col, end_col):
            marker = _clean_org_text(ws.cell(header_row + 2, c).value)
            if marker:
                labeled_cols.append(c)
        if not labeled_cols:
            continue
        fact_col = labeled_cols[0]
        plan_col = labeled_cols[1] if len(labeled_cols) > 1 else None
        specs.append({"month": month_label, "fact_col": fact_col, "plan_col": plan_col})

    total_spec = None
    if total_col:
        labeled_cols = []
        for c in range(total_col, ws.max_column + 1):
            marker = _clean_org_text(ws.cell(header_row + 2, c).value)
            if marker:
                labeled_cols.append(c)
        if labeled_cols:
            total_spec = {
                "fact_col": labeled_cols[0],
                "plan_col": labeled_cols[1] if len(labeled_cols) > 1 else None,
            }

    return specs, total_spec


def _make_month_maps(monthly_values, months):
    fact = {}
    plan = {}
    for month in months:
        x = monthly_values.get(month, {})
        fact[month] = round(float(x.get("fact_h") or 0.0), 2)
        plan[month] = round(float(x.get("plan_h") or 0.0), 2)
    return fact, plan


def build_personnel_org_payload(data_dir: Path) -> dict | None:
    """Пейлоад как раньше в personnel_org_usage.json."""
    src = find_org_personnel_analysis_xlsx(data_dir)
    if src is None:
        return None

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    sheet_title = ws.title

    header_row = _find_org_header_start(ws)
    data_start_row = header_row + 3
    month_specs, total_spec = _build_org_month_specs(ws, header_row)
    months = [x["month"] for x in month_specs]

    rows = []
    current_org = None
    current_department = None

    for r in range(data_start_row, ws.max_row + 1):
        label = _clean_org_text(ws.cell(r, 1).value)
        if not label:
            continue
        if label.lower() == "итого":
            continue

        monthly = {}
        sum_fact = 0.0
        sum_plan = 0.0
        for spec in month_specs:
            month = spec["month"]
            fact_h = _parse_hours_org(ws.cell(r, spec["fact_col"]).value) if spec.get("fact_col") else 0.0
            plan_h = _parse_hours_org(ws.cell(r, spec["plan_col"]).value) if spec.get("plan_col") else 0.0
            monthly[month] = {"fact_h": round(fact_h, 2), "plan_h": round(plan_h, 2)}
            sum_fact += fact_h
            sum_plan += plan_h

        total_fact = _parse_hours_org(ws.cell(r, total_spec["fact_col"]).value) if total_spec and total_spec.get("fact_col") else 0.0
        total_plan = _parse_hours_org(ws.cell(r, total_spec["plan_col"]).value) if total_spec and total_spec.get("plan_col") else 0.0
        fact_h = total_fact if total_fact > 0 else sum_fact
        plan_h = total_plan if total_plan > 0 else sum_plan

        if _is_fio_org(label):
            row_type = "employee"
        elif _is_org_row_org(label, current_org):
            row_type = "organization"
        else:
            row_type = "department"

        if row_type == "organization":
            current_org = label
            current_department = None
        elif row_type == "department":
            current_department = label

        utilization = (fact_h / plan_h * 100.0) if plan_h > 0 else None
        rows.append(
            {
                "type": row_type,
                "label": label,
                "organization": current_org,
                "department": current_department if row_type == "employee" else (label if row_type == "department" else None),
                "employee": label if row_type == "employee" else None,
                "fact_h": round(fact_h, 2),
                "plan_h": round(plan_h, 2),
                "utilization_pct": round(utilization, 2) if utilization is not None else None,
                "monthly": monthly,
            }
        )

    wb.close()

    employee_rows = [x for x in rows if x["type"] == "employee" and x["employee"]]
    for x in employee_rows:
        monthly_fact, monthly_plan = _make_month_maps(x.get("monthly") or {}, months)
        x["monthly_fact_h"] = monthly_fact
        x["monthly_plan_h"] = monthly_plan
        x.pop("monthly", None)
        x.pop("type", None)
        x.pop("label", None)

    employee_rows.sort(key=lambda x: x["fact_h"], reverse=True)

    dep_aggr = defaultdict(lambda: {"fact_h": 0.0, "plan_h": 0.0, "employees": 0, "monthly_fact_h": defaultdict(float), "monthly_plan_h": defaultdict(float)})
    org_aggr = defaultdict(lambda: {"fact_h": 0.0, "plan_h": 0.0, "employees": 0, "monthly_fact_h": defaultdict(float), "monthly_plan_h": defaultdict(float)})

    for e in employee_rows:
        dep_key = (e["organization"], e["department"])
        dep_aggr[dep_key]["fact_h"] += float(e["fact_h"] or 0.0)
        dep_aggr[dep_key]["plan_h"] += float(e["plan_h"] or 0.0)
        dep_aggr[dep_key]["employees"] += 1
        for month in months:
            dep_aggr[dep_key]["monthly_fact_h"][month] += float((e.get("monthly_fact_h") or {}).get(month) or 0.0)
            dep_aggr[dep_key]["monthly_plan_h"][month] += float((e.get("monthly_plan_h") or {}).get(month) or 0.0)

        org_key = e["organization"]
        org_aggr[org_key]["fact_h"] += float(e["fact_h"] or 0.0)
        org_aggr[org_key]["plan_h"] += float(e["plan_h"] or 0.0)
        org_aggr[org_key]["employees"] += 1
        for month in months:
            org_aggr[org_key]["monthly_fact_h"][month] += float((e.get("monthly_fact_h") or {}).get(month) or 0.0)
            org_aggr[org_key]["monthly_plan_h"][month] += float((e.get("monthly_plan_h") or {}).get(month) or 0.0)

    departments = []
    for (org, dep), v in dep_aggr.items():
        util = (v["fact_h"] / v["plan_h"] * 100.0) if v["plan_h"] > 0 else None
        departments.append(
            {
                "organization": org,
                "department": dep,
                "employees": int(v["employees"]),
                "fact_h": round(v["fact_h"], 2),
                "plan_h": round(v["plan_h"], 2),
                "utilization_pct": round(util, 2) if util is not None else None,
                "monthly_fact_h": {m: round(v["monthly_fact_h"][m], 2) for m in months},
                "monthly_plan_h": {m: round(v["monthly_plan_h"][m], 2) for m in months},
            }
        )
    departments.sort(key=lambda x: x["fact_h"], reverse=True)

    organizations = []
    for org, v in org_aggr.items():
        util = (v["fact_h"] / v["plan_h"] * 100.0) if v["plan_h"] > 0 else None
        organizations.append(
            {
                "organization": org,
                "employees": int(v["employees"]),
                "fact_h": round(v["fact_h"], 2),
                "plan_h": round(v["plan_h"], 2),
                "utilization_pct": round(util, 2) if util is not None else None,
                "monthly_fact_h": {m: round(v["monthly_fact_h"][m], 2) for m in months},
                "monthly_plan_h": {m: round(v["monthly_plan_h"][m], 2) for m in months},
            }
        )
    organizations.sort(key=lambda x: x["fact_h"], reverse=True)

    top_departments = departments[:12]
    return {
        "meta": {
            "source": src.name,
            "sheet": sheet_title,
            "purpose": "Personnel usage by organizations/departments/employees for dashboard and AI agent",
            "months": months,
            "organizations_count": len(organizations),
            "departments_count": len(departments),
            "employees_count": len(employee_rows),
        },
        "table": {
            "title": "Использование персонала по организациям и подразделениям",
            "columns": [
                {"key": "organization", "label": "Организация"},
                {"key": "department", "label": "Подразделение"},
                {"key": "employee", "label": "Сотрудник"},
                {"key": "fact_h", "label": "Факт, ч"},
                {"key": "plan_h", "label": "План, ч"},
                {"key": "utilization_pct", "label": "Выполнение, %"},
            ],
            "rows": employee_rows,
        },
        "departments": departments,
        "organizations": organizations,
        "chart": {
            "type": "chart",
            "title": "Топ подразделений по фактическим трудозатратам (ч)",
            "chartType": "bar",
            "categories": [f"{x['department']} · {x['organization']}" for x in top_departments],
            "series": [
                {"name": "Факт, ч", "data": [x["fact_h"] for x in top_departments]},
                {"name": "План, ч", "data": [x["plan_h"] for x in top_departments]},
            ],
        },
    }
